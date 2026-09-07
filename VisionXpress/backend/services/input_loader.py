#Converse use slip little manual planning time reduction or reduced on time, same show, userrattle specifically notesport, athlete meth corner laptop mere fourth generation, movement, click show wiswer click zero training, dress plea slot uper nice symbols, you move symbols related to paiga, so relief drink, fighting captain, luft college ninety two already chart displayed low no mithelation then change thisstuff past it onimport csv
from pathlib import Path

from openpyxl import load_workbook

import csv
def load_file(file_path: str) -> list[dict]:
    """
    Load a CSV or XLSX file and return its rows as dictionaries.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = path.suffix.lower()

    if suffix == ".csv":
        return _load_csv(path)

    if suffix in [".xlsx", ".xlsm"]:
        return _load_xlsx(path)

    raise ValueError(
        f"Unsupported file type: {suffix}. "
        "Use CSV or XLSX."
    )


def _load_csv(path: Path) -> list[dict]:
    """
    Load CSV file.
    """

    with path.open(
        mode="r",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        reader = csv.DictReader(file)

        return [
            _clean_row(row)
            for row in reader
        ]


def _load_xlsx(path: Path) -> list[dict]:
    """
    Load the first worksheet from an XLSX file.
    """

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    rows = worksheet.iter_rows(
        values_only=True
    )

    headers = next(rows)

    headers = [
        str(header).strip()
        if header is not None
        else ""
        for header in headers
    ]

    data = []

    for row in rows:

        record = {}

        for header, value in zip(headers, row):

            if header:
                record[header] = value

        data.append(_clean_row(record))

    workbook.close()

    return data


def _clean_row(row: dict) -> dict:
    """
    Clean whitespace from field names and string values.
    """

    cleaned = {}

    for key, value in row.items():

        clean_key = str(key).strip()

        if isinstance(value, str):
            value = value.strip()

        cleaned[clean_key] = value

    return cleaned


def load_maintenance_file(file_path: str) -> list[dict]:
    """
    Load one maintenance department file.

    TMS, TDMS and SMMS all use the same normalized
    maintenance-request structure.
    """

    return load_file(file_path)


def load_coa_file(file_path: str) -> list[dict]:
    """
    Load COA train schedule data.
    """

    return load_file(file_path)


def load_all_maintenance(
    tms_path: str,
    tdms_path: str,
    smms_path: str,
) -> list[dict]:
    """
    Load maintenance requests from TMS, TDMS and SMMS
    into one combined list.
    """

    tasks = []

    tasks.extend(
        load_maintenance_file(tms_path)
    )

    tasks.extend(
        load_maintenance_file(tdms_path)
    )

    tasks.extend(
        load_maintenance_file(smms_path)
    )

    return tasks